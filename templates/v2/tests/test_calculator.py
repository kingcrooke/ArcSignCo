#!/usr/bin/env python3
"""Calculator formula + LibreOffice recalc assertions."""
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output"
sys.path.insert(0, str(ROOT / "scripts"))

from generate_arc_calculator_v2 import ITEMS
from libreoffice_recalc import amounts_by_label, recalc_xlsx_to_values


def row_map(ws) -> dict[str, int]:
    m = {}
    for row in range(1, ws.max_row + 1):
        lab = ws.cell(row, 5).value
        if not lab:
            continue
        key = {
            "Subtotal": "sub",
            "Sales tax": "tax",
            "TOTAL": "total",
            "Deposit (50%) - before work starts": "dep",
            "Production (25%) - before install": "prod",
            "Additional costs": "addl",
            "Final (25%) + additional costs - Net 30": "final",
            "Card processing fee": "fee",
            "TOTAL BILLED (incl. extras / card fee)": "grand",
            "Amount paid": "paid",
            "BALANCE DUE": "bal",
            "Tax exempt? (Yes/No)": "exempt",
            "Paying by card? (Yes/No)": "card",
        }.get(str(lab))
        if key:
            m[key] = row
    return m


def first_item_row(ws) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == "01":
            return row
    raise RuntimeError("line item row not found")


def approx(a, b, tol=0.01):
    return abs(float(a) - float(b)) <= tol


def configure_sheet(ws, *, subtotal_qty_rate=(1, 3300), addl=0, paid=0, exempt="No", card="No"):
    first = first_item_row(ws)
    for i in range(len(ITEMS)):
        row = first + i
        ws.cell(row, 4, 0)
        ws.cell(row, 6, 0)
    ws.cell(first, 4, subtotal_qty_rate[0])
    ws.cell(first, 6, subtotal_qty_rate[1])
    L = row_map(ws)
    ws.cell(L["addl"], 7, addl)
    ws.cell(L["paid"], 7, paid)
    ws.cell(L["exempt"], 7, exempt)
    ws.cell(L["card"], 7, card)
    return L


def assert_formulas(ws):
    first = first_item_row(ws)
    for row in range(first, first + len(ITEMS)):
        assert str(ws.cell(row, 7).value).startswith("=")
    L = row_map(ws)
    for key in ("sub", "tax", "total", "dep", "prod", "final", "fee", "grand", "bal"):
        v = ws.cell(L[key], 7).value
        if not (isinstance(v, str) and v.startswith("=")):
            raise AssertionError(f"{key} lost formula: {v}")


def main() -> int:
    src = OUT / "Arc_Estimate_Invoice_Calculator_v2.xlsx"
    if not src.is_file():
        print(f"FAIL: missing {src}")
        return 1

    wb_check = load_workbook(src, data_only=False)
    assert_formulas(wb_check["Estimate"])
    assert_formulas(wb_check["Invoice"])

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "scenario1.xlsx"
        wb = load_workbook(src)
        ws = wb["Estimate"]
        L = configure_sheet(ws, addl=100, paid=1000, exempt="No", card="No")
        wb.save(path)
        grid = recalc_xlsx_to_values(path, 0)
        amts = amounts_by_label(grid)
        for name, got, want in (
            ("tax", amts.get("Sales tax"), 292.88),
            ("total", amts.get("TOTAL"), 3592.88),
            ("deposit", amts.get("Deposit (50%) - before work starts"), 1796.44),
            ("production", amts.get("Production (25%) - before install"), 898.22),
            ("final", amts.get("Final (25%) + additional costs - Net 30"), 998.22),
            ("balance", amts.get("BALANCE DUE"), 2692.88),
        ):
            if not approx(got, want):
                print(f"FAIL scenario 1 {name}: got {got}, want {want}")
                return 1
        print("PASS scenario 1 (taxed, no card fee)")

        path2 = Path(td) / "scenario2.xlsx"
        wb2 = load_workbook(src)
        ws2 = wb2["Estimate"]
        L2 = configure_sheet(ws2, paid=1650, exempt="Yes", card="Yes")
        wb2.save(path2)
        grid2 = recalc_xlsx_to_values(path2, 0)
        amts2 = amounts_by_label(grid2)
        for name, got, want in (
            ("total", amts2.get("TOTAL"), 3300),
            ("card fee", amts2.get("Card processing fee"), 115.50),
            ("balance", amts2.get("BALANCE DUE"), 1765.50),
        ):
            if not approx(got, want):
                print(f"FAIL scenario 2 {name}: got {got}, want {want}")
                return 1
        print("PASS scenario 2 (tax exempt + card fee)")

    print("PASS: formulas still present on Estimate and Invoice sheets")
    print("PASS: test_calculator.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
