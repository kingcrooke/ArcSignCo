#!/usr/bin/env python3
"""Edit DOCX placeholders, re-render PDF, verify content and single-page layout."""
import shutil
import sys
import tempfile
from pathlib import Path

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output"
sys.path.insert(0, str(ROOT / "scripts"))
from convert_docx import convert_docx


def replace_in_docx(src: Path, dst: Path, replacements: dict[str, str]) -> None:
    shutil.copy(src, dst)
    doc = Document(dst)
    for p in doc.paragraphs:
        for old, new in replacements.items():
            if old in p.text:
                for run in p.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old, new in replacements.items():
                        if old in p.text:
                            for run in p.runs:
                                if old in run.text:
                                    run.text = run.text.replace(old, new)
                for nt in cell.tables:
                    for r2 in nt.rows:
                        for c2 in r2.cells:
                            for p in c2.paragraphs:
                                for old, new in replacements.items():
                                    if old in p.text:
                                        for run in p.runs:
                                            if old in run.text:
                                                run.text = run.text.replace(old, new)
    doc.save(dst)


def main() -> int:
    cases = [
        (OUT / "Arc_Estimate_Template_v2.docx", "estimate"),
        (OUT / "Arc_Invoice_Template_v2.docx", "invoice"),
    ]
    for src, label in cases:
        if not src.is_file():
            print(f"FAIL: missing {src}")
            return 1
        with tempfile.TemporaryDirectory() as td:
            td_path = Path(td)
            edited = td_path / src.name
            replace_in_docx(
                src,
                edited,
                {"[CLIENT NAME]": "Test Client LLC", "[AMOUNT]": "$1,234.56"},
            )
            pdf, _png = convert_docx(edited, td_path)
            text = "\n".join(page.get_text() for page in pymupdf.open(pdf))
            pages = len(pymupdf.open(pdf))
            if "Test Client LLC" not in text:
                print(f"FAIL: {label} edit — client name not in PDF")
                return 1
            if "$1,234.56" not in text and "1,234.56" not in text:
                print(f"FAIL: {label} edit — amount not in PDF")
                return 1
            if pages != 1:
                print(f"FAIL: {label} edit — page count {pages}, expected 1")
                return 1
            print(f"PASS: {label} edit — client + amount in PDF, 1 page")
    print("PASS: test_docx_edit.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
