#!/usr/bin/env python3
"""Arc Signage Co LLC - Estimate / Invoice calculator v2 (openpyxl, live formulas)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from lib_paths import OUT

OUT.mkdir(parents=True, exist_ok=True)
XLSX = OUT / "Arc_Estimate_Invoice_Calculator_v2.xlsx"
NAVY, GOLD, INPUT = "0B1D33", "D4A843", "FFF7E6"
F = "Arial"
thin = Side(style="thin", color="D9DDE3")
money = '"$"#,##0.00'
ITEMS = [
    "Design",
    "Engineering drawings",
    "Fabrication",
    "Installation",
    "Permit coordination",
    "[ADDITIONAL ITEM]",
    "[ADDITIONAL ITEM]",
]


def sheet(ws, est):
    ws.sheet_view.showGridLines = False
    widths = {"A": 5, "B": 24, "C": 46, "D": 9, "E": 9, "F": 13, "G": 15}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws["A1"] = "Arc Signage Co LLC"
    ws["A1"].font = Font(name=F, size=14, bold=True, color=NAVY)
    ws["A2"] = (
        "1974 Crotona Ave, Bronx, NY 10457 · (347) 450-2110 (call or text) · "
        "jc@arcsignco.com · arcsignco.com"
    )
    ws["A2"].font = Font(name=F, size=9, color="6B7280")
    ws["G1"] = "ESTIMATE / PROPOSAL" if est else "INVOICE"
    ws["G1"].font = Font(name=F, size=14, bold=True, color=NAVY)
    ws["G1"].alignment = Alignment(horizontal="right")
    meta = (
        [
            ("Estimate No.", "ARC-[YYYY]-[MMDD]-[XXX]"),
            ("Date", "[DATE]"),
            ("Valid until", "[DATE + 30 DAYS]"),
        ]
        if est
        else [
            ("Invoice No.", "INV-[YYYY]-[MMDD]-[XXX]"),
            ("Invoice date", "[DATE]"),
            ("Due date", "[DUE DATE]"),
            ("Estimate ref.", "ARC-[YYYY]-[MMDD]-[XXX]"),
        ]
    )
    meta += [("Client", "[CLIENT NAME]"), ("Project", "[PROJECT NAME]")]
    r = 4
    for k, v in meta:
        ws.cell(r, 2, k).font = Font(name=F, size=9, bold=True, color="6B7280")
        c = ws.cell(r, 3, v)
        c.font = Font(name=F, size=9, bold=True, color="B45F06")
        c.fill = PatternFill("solid", fgColor=INPUT)
        r += 1
    hr = r + 1
    for i, h in enumerate(["#", "Item", "Description", "Qty", "Unit", "Rate", "Amount"], start=1):
        c = ws.cell(hr, i, h)
        c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="right" if i >= 4 else "left")
    first = hr + 1
    for n, item in enumerate(ITEMS):
        row = first + n
        ws.cell(row, 1, f"{n + 1:02d}").font = Font(name=F, size=9, bold=True, color=GOLD)
        ws.cell(row, 2, item).font = Font(name=F, size=9, bold=True, color=NAVY)
        ws.cell(row, 3, "[DESCRIPTION]").font = Font(name=F, size=9, color="6B7280")
        ws.cell(row, 4, 0)
        ws.cell(row, 5, "[UNIT]")
        ws.cell(row, 6, 0)
        ws.cell(row, 7, f"=D{row}*F{row}")
        for col in (4, 5, 6):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=INPUT)
        ws.cell(row, 6).number_format = money
        ws.cell(row, 7).number_format = money
        for col in range(1, 8):
            ws.cell(row, col).border = Border(bottom=thin)
            if ws.cell(row, col).font.name != F:
                ws.cell(row, col).font = Font(name=F, size=9)
    last = first + len(ITEMS) - 1
    t = last + 2
    L = {}

    def line(key, label, formula, pct=None, fmt=money, inp=False, strong=False, note=None):
        nonlocal t
        ws.cell(t, 5, label).font = Font(
            name=F, size=9, bold=strong, color="FFFFFF" if strong else "1F2937"
        )
        if pct is not None:
            c = ws.cell(t, 6, pct)
            c.number_format = "0.000%" if pct and pct < 0.1 else "0%"
            c.fill = PatternFill("solid", fgColor=INPUT)
        c = ws.cell(t, 7, formula)
        c.number_format = fmt
        c.font = Font(name=F, size=9, bold=strong, color=GOLD if strong else "1F2937")
        if inp:
            c.fill = PatternFill("solid", fgColor=INPUT)
        if strong:
            for col in (5, 6, 7):
                ws.cell(t, col).fill = PatternFill("solid", fgColor=NAVY)
        if note:
            ws.cell(t, 3, note).font = Font(name=F, size=8, italic=True, color="6B7280")
        L[key] = t
        t += 1

    line("sub", "Subtotal", f"=SUM(G{first}:G{last})")
    line("rate", "Sales tax rate", "", pct=0.08875, note="Editable rate (8.875%)")
    ws.cell(L["rate"], 7).value = None
    line(
        "exempt",
        "Tax exempt? (Yes/No)",
        "No",
        fmt="@",
        inp=True,
        note="Yes only with a valid exemption or resale certificate on file before invoicing",
    )
    line("tax", "Sales tax", f'=IF(G{L["exempt"]}="Yes",0,ROUND(G{L["sub"]}*F{L["rate"]},2))')
    line("total", "TOTAL", f'=G{L["sub"]}+G{L["tax"]}', strong=True)
    t += 1
    dep_row = t
    line("dep", "Deposit (50%) - before work starts", f'=ROUND(G{L["total"]}*F{dep_row},2)', pct=0.5)
    prod_row = t
    line("prod", "Production (25%) - before install", f'=ROUND(G{L["total"]}*F{prod_row},2)', pct=0.25)
    line(
        "addl",
        "Additional costs",
        0,
        inp=True,
        note="Enter any change orders / extras billed with the final payment",
    )
    final_row = t
    line(
        "final",
        "Final (25%) + additional costs - Net 30",
        f'=G{L["total"]}-G{L["dep"]}-G{L["prod"]}+G{L["addl"]}',
        pct=0.25,
        note="Final = total - deposit - production + additional costs (absorbs rounding)",
    )
    t += 1
    card_row = t
    line("card", "Paying by card? (Yes/No)", "No", fmt="@", inp=True)
    fee_row = t
    line(
        "fee",
        "Card processing fee",
        f'=IF(G{card_row}="Yes",ROUND((G{L["total"]}+G{L["addl"]})*F{fee_row},2),0)',
        pct=0.035,
        note="3.5% of total + additional costs when paid by card",
    )
    line("grand", "TOTAL BILLED (incl. extras / card fee)", f'=G{L["total"]}+G{L["addl"]}+G{L["fee"]}', strong=True)
    line("paid", "Amount paid", 0, inp=True)
    line("bal", "BALANCE DUE", f'=G{L["grand"]}-G{L["paid"]}', strong=True)
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"G{L['exempt']}")
    dv.add(f"G{L['card']}")
    ws.cell(t + 1, 2, "Yellow cells are inputs. All other amounts are live formulas.").font = Font(
        name=F, size=8, italic=True, color="6B7280"
    )
    ws.print_area = f"A1:G{t + 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return L, first


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Estimate"
    sheet(ws, True)
    sheet(wb.create_sheet("Invoice"), False)
    return wb


if __name__ == "__main__":
    wb = build_workbook()
    wb.save(XLSX)
    print(XLSX)
